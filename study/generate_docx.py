import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    return p

def format_inline_runs(paragraph, text, base_font_size=10.5, is_italic=False):
    # Regex to split on bold (**...**), inline code (`...`), and regular text
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.size = Pt(base_font_size)
        run.font.name = "Segoe UI"
        
        if token.startswith("**") and token.endswith("**"):
            run.text = token[2:-2]
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith("`") and token.endswith("`"):
            run.text = f" {token[1:-1]} "
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(194, 65, 12)
        elif token.startswith("*") and token.endswith("*"):
            run.text = token[1:-1]
            run.italic = True
        else:
            run.text = token
            run.font.color.rgb = RGBColor(51, 65, 85)
            if is_italic:
                run.italic = True

def create_code_block(doc, code_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    
    text = "\n".join(code_lines)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
    # Title Header Block
    title_p = add_styled_paragraph(doc, "", space_before=0, space_after=2)
    r1 = title_p.add_run("Vibe Editor: Master Architecture & Technical Interview Defense Guide")
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(22)
    r1.bold = True
    r1.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    
    sub_p = add_styled_paragraph(doc, "", space_before=2, space_after=16)
    r2 = sub_p.add_run("Comprehensive Technical Documentation, System Design, Edge Cases & Defense Blueprint")
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(11)
    r2.italic = True
    r2.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
    
    # Divider
    div = doc.add_paragraph()
    div.paragraph_format.space_after = Pt(12)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="2563EB"/></w:pBdr>')
    div._p.get_or_add_pPr().append(pBdr)
    
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.startswith("```"):
            if in_code_block:
                create_code_block(doc, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Blank line
        if not line.strip():
            i += 1
            continue
            
        # Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
            
        # Markdown Headings
        if line.startswith("# "):
            # We already have title, treat as Section Title
            heading_text = line[2:].strip()
            p = add_styled_paragraph(doc, "", space_before=16, space_after=6)
            r = p.add_run(heading_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(17)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
            i += 1
            continue
            
        if line.startswith("## "):
            heading_text = line[3:].strip()
            p = add_styled_paragraph(doc, "", space_before=14, space_after=5)
            r = p.add_run(heading_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(14)
            r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
            
        if line.startswith("### "):
            heading_text = line[4:].strip()
            p = add_styled_paragraph(doc, "", space_before=10, space_after=3)
            r = p.add_run(heading_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(12)
            r.bold = True
            r.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue
            i += 1
            continue
            
        if line.startswith("#### "):
            heading_text = line[5:].strip()
            p = add_styled_paragraph(doc, "", space_before=8, space_after=2)
            r = p.add_run(heading_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue
            
        # Blockquote ("> ...")
        if line.startswith(">"):
            quote_text = line[1:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="none"/>'
                f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, quote_text, base_font_size=10, is_italic=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
            i += 1
            continue
            
        # Tables ("| ... |")
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Filter out separator lines like |:---|:---|
            data_rows = [row for row in table_lines if not re.match(r'^\|(\s*:?-+:?\s*\|)+$', row)]
            if data_rows:
                header_cols = [c.strip() for c in data_rows[0].strip('|').split('|')]
                num_cols = len(header_cols)
                
                table = doc.add_table(rows=len(data_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Header styling
                for col_idx, col_text in enumerate(header_cols):
                    cell = table.cell(0, col_idx)
                    set_cell_background(cell, "1E293B") # Dark Slate Header
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    run = p.add_run(col_text.replace('**', ''))
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    
                # Data rows styling
                for row_idx, row_str in enumerate(data_rows[1:], start=1):
                    cols = [c.strip() for c in row_str.strip('|').split('|')]
                    # Handle cell count mismatch if any
                    for col_idx in range(num_cols):
                        val = cols[col_idx] if col_idx < len(cols) else ""
                        cell = table.cell(row_idx, col_idx)
                        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        
                        # Add thin border
                        tcPr = cell._tc.get_or_add_tcPr()
                        borders = parse_xml(
                            f'<w:tcBorders {nsdecls("w")}>'
                            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                            f'</w:tcBorders>'
                        )
                        tcPr.append(borders)
                        
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        # Replace <br> with newline for formatting
                        val = val.replace('<br>', '\n')
                        format_inline_runs(p, val, base_font_size=9)
                        
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue
            
        # Bullet list items
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            content = line.strip()[2:].strip()
            indent_level = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            format_inline_runs(p, content, base_font_size=10.5)
            i += 1
            continue
            
        # Numbered list items
        m_num = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if m_num:
            num = m_num.group(1)
            content = m_num.group(2)
            indent_level = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            format_inline_runs(p, content, base_font_size=10.5)
            i += 1
            continue
            
        # Regular paragraph
        p = add_styled_paragraph(doc, "", space_before=2, space_after=5)
        format_inline_runs(p, line, base_font_size=10.5)
        i += 1
        
    doc.save(docx_path)
    print(f"Successfully generated Word document at: {docx_path}")

if __name__ == "__main__":
    src_md = r"d:\Project\Vibe Editor\my-app\study\master_interview_architecture_guide.md"
    out_docx = r"d:\Project\Vibe Editor\my-app\study\Vibe_Editor_Master_Architecture_Guide.docx"
    parse_markdown_to_docx(src_md, out_docx)
